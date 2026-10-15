#!/usr/bin/env python3
"""(不要改名成 inspect.py:会遮住 Python 标准库的 inspect 模块,numpy/openpyxl 一导入就崩。)

打印 .docx / .xlsx / .pptx 的结构摘要,改文件前后各跑一次核对。

用法: python3 office_inspect.py <文件> [--max-rows N]
只读,不修改文件。每段文字截到 120 字,避免把整份文档塞进对话。
"""
import os
import sys

MAX_TEXT = 120


def clip(text):
    text = " ".join(str(text).split())
    return text if len(text) <= MAX_TEXT else text[:MAX_TEXT] + "…"


def inspect_docx(path):
    from docx import Document

    doc = Document(path)
    paras = [p for p in doc.paragraphs if p.text.strip()]
    print(f"Word: {len(doc.paragraphs)} 段(非空 {len(paras)}), {len(doc.tables)} 个表格, {len(doc.sections)} 节")
    for i, p in enumerate(paras[:60]):
        print(f"  [{i}] ({p.style.name}, {len(p.runs)} runs) {clip(p.text)}")
    if len(paras) > 60:
        print(f"  … 还有 {len(paras) - 60} 段")
    for t, table in enumerate(doc.tables[:10]):
        print(f"  表格 {t}: {len(table.rows)} 行 × {len(table.columns)} 列")
        for row in table.rows[:3]:
            print("    | " + " | ".join(clip(c.text) for c in row.cells))


def inspect_xlsx(path, max_rows):
    from openpyxl import load_workbook

    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True, read_only=True)
    print(f"Excel: {len(formulas.worksheets)} 个工作表")
    for ws in formulas.worksheets:
        count = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        print(f"  工作表 「{ws.title}」: {ws.max_row} 行 × {ws.max_column} 列, 公式 {count} 个")
        vws = values[ws.title]
        for r, row in enumerate(vws.iter_rows(values_only=True)):
            if r >= max_rows:
                print(f"    … 只显示前 {max_rows} 行")
                break
            print("    | " + " | ".join("" if v is None else clip(v) for v in row))


def inspect_pptx(path):
    from pptx import Presentation

    prs = Presentation(path)
    print(f"PowerPoint: {len(prs.slides)} 页")
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(clip(shape.text_frame.text))
        title = texts[0] if texts else "(无文字)"
        print(f"  第 {i + 1} 页: {title}")
        for t in texts[1:4]:
            print(f"      {t}")


def main():
    if len(sys.argv) < 2:
        sys.exit("用法: python3 office_inspect.py <文件> [--max-rows N]")
    path = sys.argv[1]
    max_rows = int(sys.argv[sys.argv.index("--max-rows") + 1]) if "--max-rows" in sys.argv else 8
    if not os.path.isfile(path):
        sys.exit(f"文件不存在: {path}")
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".docx":
            inspect_docx(path)
        elif ext == ".xlsx" or ext == ".xlsm":
            inspect_xlsx(path, max_rows)
        elif ext == ".pptx":
            inspect_pptx(path)
        else:
            sys.exit(f"不支持的格式 {ext}:只处理 .docx / .xlsx / .pptx(旧格式 .doc/.xls/.ppt 请用户另存为新格式)")
    except ImportError as e:
        sys.exit(f"缺少依赖({e.name or e}),先按 SKILL.md 安装")


if __name__ == "__main__":
    main()
